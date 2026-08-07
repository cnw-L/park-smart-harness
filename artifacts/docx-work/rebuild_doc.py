from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


ROOT = Path(r"D:\park-smart-harness\artifacts\docx-work")
MEDIA = ROOT / "unpacked" / "word" / "media"
OUTPUT = ROOT / "状况-排版优化版.docx"


CAPTIONS = {
    1: "系统通过对话快速汇总设备可靠率、可用率及故障率，并提示需要重点关注的异常指标。",
    2: "设备健康总览集中展示故障数量、可靠率、可用率及平均维修时长等核心运行指标。",
    3: "系统默认统计近7天工单数量、类型、完成率和超时情况，并可继续下钻查询具体明细。",
    4: "系统列出近7天生成的工单及其当前状态、创建时间，便于快速定位待处理事项。",
    5: "针对空调机组，可直接查询近7天报修工单数量、处理状态及工单生成时间范围。",
    6: "系统支持按自然语言查询本月、昨日等时间范围的工单概况，并准确识别对应日期。",
    7: "月度统计图展示工单类型分布以及待调度、待指派、已关闭等各状态数量。",
    8: "近7天统计图同步呈现工单总量、类型占比和各处理状态，便于核对对话查询结果。",
    9: "单日统计结果显示指定日期的工单数量及状态分布，可用于验证昨日工单查询。",
    10: "工单列表按编号、标题、类型、生成时间和状态展示查询结果，支持查看与调度操作。",
    11: "系统可根据自然语言生成报修工单草稿，待用户确认后创建并返回工单编号与状态。",
    12: "在执行创建维修工单前，系统弹出确认窗口，明确目标设备、操作内容和风险级别。",
    13: "卡片式工单列表支持按类型、状态和时间筛选，并直观展示工单耗时与紧急程度。",
    14: "工单详情页完整记录报修信息、事件内容、报修时间、联系人及当前处理状态。",
    15: "系统查询设备实时状态后，仅在发现异常时建议创建报修工单，避免重复或无效派单。",
    16: "AI助手与设备档案联动，可在同一页面查询设备状态并继续执行相关运维操作。",
    17: "通过远程摄像头画面可查看现场环境，并使用云台控制辅助核实设备及周边状况。",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for style_name, size, color in (
        ("Title", 22, "17365D"),
        ("Heading 1", 16, "17365D"),
        ("Heading 2", 13, "2F5597"),
    ):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(8)
    if level == 1:
        p_pr = paragraph._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "5")
        bottom.set(qn("w:color"), "5B9BD5")
        border.append(bottom)
        p_pr.append(border)


def add_figure(doc, number):
    path = MEDIA / f"image{number}.png"
    with Image.open(path) as image:
        width_px, height_px = image.size
    max_width_cm = 16.0
    max_height_cm = 15.8
    width_cm = max_width_cm
    height_cm = width_cm * height_px / width_px
    if height_cm > max_height_cm:
        height_cm = max_height_cm
        width_cm = height_cm * width_px / height_px

    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.keep_with_next = True
    image_p.paragraph_format.space_before = Pt(4)
    image_p.paragraph_format.space_after = Pt(4)
    image_p.add_run().add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_after = Pt(12)
    run = caption.add_run(f"图 {number}  {CAPTIONS[number]}")
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(89, 89, 89)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)
configure_styles(doc)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
title.add_run("智慧园区系统功能状况说明")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
subtitle_run = subtitle.add_run("设备状态与工单管理功能截图说明")
subtitle_run.font.name = "微软雅黑"
subtitle_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
subtitle_run.font.size = Pt(11)
subtitle_run.font.color.rgb = RGBColor(117, 117, 117)

intro = doc.add_table(rows=1, cols=1)
intro.autofit = False
intro.columns[0].width = Cm(16.0)
cell = intro.cell(0, 0)
cell.width = Cm(16.0)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
set_cell_shading(cell, "EAF2F8")
cell.text = "说明：以下截图按“设备概况—工单查询—工单执行—设备联动”的业务顺序整理；未指定查询时间时，工单统计默认采用近7天范围。"
for p in cell.paragraphs:
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.right_indent = Cm(0.2)

add_heading(doc, "一、设备状态概况查询", 1)
add_figure(doc, 1)
add_figure(doc, 2)

add_heading(doc, "二、工单情况查询", 1)
add_heading(doc, "2.1 对话查询与明细下钻", 2)
for n in range(3, 6):
    add_figure(doc, n)

add_heading(doc, "2.2 多时间范围统计核对", 2)
for n in range(6, 10):
    add_figure(doc, n)

add_heading(doc, "2.3 工单列表、创建与确认", 2)
for n in range(10, 15):
    add_figure(doc, n)

add_heading(doc, "三、设备状态联动与远程查看", 1)
for n in range(15, 18):
    add_figure(doc, n)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("智慧园区系统功能状况说明")
footer_run.font.name = "微软雅黑"
footer_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
footer_run.font.size = Pt(9)
footer_run.font.color.rgb = RGBColor(128, 128, 128)

doc.core_properties.title = "智慧园区系统功能状况说明"
doc.core_properties.subject = "设备状态与工单管理功能截图说明"
doc.core_properties.author = "Codex"
doc.save(OUTPUT)
print(OUTPUT)
