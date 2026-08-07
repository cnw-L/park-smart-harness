from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(r"D:\park-smart-harness")
SOURCE = ROOT / "artifacts" / "状况-最终整理版.docx"
OUTPUT = ROOT / "artifacts" / "状况-最终整理版（含中文工具名）.docx"

ROWS = [
    ("参数重复识别", "全部 SchemaCapability", "所有结构化工具", "一次调用会重复识别参数，可能增加响应时间，并出现前后结果不一致。", "优先使用模型已填写的参数，缺失时再补充识别。", "中"),
    ("时间格式不统一", "alarm.query\nenergy.query\nevent.query\nwork_order.query\nmeeting.query\nmeeting.reserve", "告警查询\n能耗查询\n事件查询\n工单查询\n会议室查询\n会议室预订", "“今天”“近7天”等时间在不同工具中的接收格式不一致。", "统一时间格式，并保留自然语言时间的自动转换。", "高"),
    ("部分操作依赖上一轮", "device.control_ticket.prepare\nwork_order.create_prepare\nalarm.create_order\nmeeting.reserve", "设备控制票据准备\n工单创建准备\n告警建单\n会议室预订", "缺少上一轮查询结果或必要信息时，无法直接执行。", "缺少信息时明确提示；合适的场景可先自动查询，再让用户确认。", "高"),
    ("确认流程较独立", "device.control_ticket.prepare\nwork_order.create_prepare\nevent.create_prepare\nalarm.create_order", "设备控制票据准备\n工单创建准备\n事件上报准备\n告警建单", "准备、确认和执行分为不同步骤，状态管理较复杂。", "保留确认机制，同时让待确认内容和执行结果更清楚。", "中"),
    ("设备名称容易误匹配", "device.status.query\nui.focus_camera\nui.open_device_panel", "设备状态查询\n摄像头聚焦\n打开设备面板", "设备名称中的数字、空格或房间号可能被误当成设备ID。", "清理多余空格，区分设备名称和设备ID，并增加匹配校验。", "高"),
    ("事件标题不够干净", "event.create_prepare", "事件上报准备", "“我要上报一个事件”等对话前缀会进入标题和描述。", "创建草稿前清理口语前缀，只保留真正的事件内容。", "立即"),
]

HEADERS = ["问题", "工具能力", "中文工具名", "主要表现", "优化方式", "优先级"]
WIDTHS = [Cm(2.3), Cm(4.6), Cm(4.0), Cm(5.2), Cm(6.8), Cm(1.6)]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(cell, value=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for tag in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def keep_row(row, repeat=False):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_text(cell, text, *, bold=False, color="1F1F1F", align=WD_ALIGN_PARAGRAPH.LEFT, size=7.8):
    cell.text = ""
    lines = str(text).split("\n")
    for index, line in enumerate(lines):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1 if index < len(lines) - 1 else 0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line)
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


doc = Document(SOURCE)

section = doc.sections[2]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Mm(297)
section.page_height = Mm(210)
section.top_margin = Cm(1.2)
section.bottom_margin = Cm(1.2)
section.left_margin = Cm(1.2)
section.right_margin = Cm(1.2)

old_table = next(
    table for table in doc.tables
    if table.rows and [cell.text for cell in table.rows[0].cells][:2] == ["问题", "影响范围"]
)

new_table = doc.add_table(rows=1, cols=6)
new_table.alignment = WD_TABLE_ALIGNMENT.CENTER
new_table.autofit = False
new_table.style = "Table Grid"

header = new_table.rows[0]
keep_row(header, repeat=True)
for cell, text, width in zip(header.cells, HEADERS, WIDTHS):
    cell.width = width
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    margins(cell)
    shade(cell, "2F5597")
    set_text(cell, text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, size=8.2)

for index, values in enumerate(ROWS):
    row = new_table.add_row()
    keep_row(row)
    for col, (cell, text, width) in enumerate(zip(row.cells, values, WIDTHS)):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        margins(cell)
        fill = "F7F9FC" if index % 2 else "FFFFFF"
        if col == 2:
            fill = "EAF2F8"
        if col == 5:
            fill = "FCE4D6" if text in ("高", "立即") else "FFF2CC"
        shade(cell, fill)
        color = "C65911" if col == 5 and text in ("高", "立即") else "1F1F1F"
        set_text(
            cell,
            text,
            bold=col in (0, 2, 5),
            color=color,
            align=WD_ALIGN_PARAGRAPH.CENTER if col == 5 else WD_ALIGN_PARAGRAPH.LEFT,
        )

old_table._tbl.addprevious(new_table._tbl)
old_table._element.getparent().remove(old_table._element)

doc.save(OUTPUT)
print(OUTPUT)
