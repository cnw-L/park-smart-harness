from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(r"D:\park-smart-harness")
SOURCE = ROOT / "artifacts" / "状况-排版优化版.docx"
OUTPUT = ROOT / "artifacts" / "状况-排版优化及能力验证版.docx"

ROWS = [
    (1, "device.status.query", "真实·设备网关", "✓", "空调机组106停止；送风湿度26%、送风温度30℃、新风阀开度12%。", "真数据；单设备精确查询正常。"),
    (2, "device.control_ticket.prepare", "真实·设备网关", "✓", "生成“关闭空调机组106”操作票据，状态为 needs_confirmation。", "确认门控正确，尚未执行；注册名为 device.control_ticket.prepare，并非 device.control。"),
    (3, "device.health.query", "真实·健康网关", "✓", "机组106可靠性99.79%、可用率96.77%、故障71起、故障率4.63%。", "真数据。"),
    (4, "energy.query", "真实·能耗网关", "✓", "返回“本月暂无能耗数据”。", "真实调用成功，但能耗系统在该时间窗口内无数据。"),
    (5, "work_order.query", "真实·工单网关", "✓", "近一月15单：物业报修14单、巡检1单；待调度13单、超时1单。", "真数据。"),
    (6, "work_order.create_prepare", "真实·工单网关", "✓", "生成机组106维修工单草稿，状态为 needs_confirmation。", "确认门控正确。"),
    (7, "alarm.query", "真实·告警网关", "✓", "近7天0条告警，历史告警总数2098条。", "默认7天查询窗口为空，与既有观察一致。"),
    (8, "alarm.create_order", "真实·告警网关", "✗", "请求被路由至 alarm.query；近7天无门禁告警，因此无法创建工单。", "未触发：需先获得告警候选，再发起建单，属于“查询→建单”的两轮流程。"),
    (9, "event.query", "真实·事件网关", "✓", "近7天共21条事件；3楼卫生间漏水频发，7月11日当天发生8次。", "真数据。"),
    (10, "event.create_prepare", "真实·事件网关", "✓", "生成事件上报草稿，状态为 needs_confirmation。", "待修复：标题未清洗，“我要上报一个事件：”前缀泄漏至标题和描述。"),
    (11, "ui.navigate", "真实·菜单网关", "✓", "解析门禁管理页面并发出 navigate_route 指令。", "菜单树为真实解析结果。"),
    (12, "knowledge.qa", "真实·Milvus+rerank", "✓", "检索5条知识后，诚实拒答“未收录消防巡检流程”。", "接地良好，无内容臆造；QA严格化策略生效。"),
    (13, "ui.focus_camera", "前端指令·无后端", "✓", "发出 focus_camera_view 指令，目标为1号摄像头。", "纯UI指令。"),
    (14, "ui.open_device_panel", "前端指令·无后端", "✓", "发出 open_device_panel 指令，目标为机组106。", "纯UI指令。"),
    (15, "date.current", "本地·无后端", "✓", "返回2026-07-14 11:02，星期二。", "本地确定性结果。"),
    (16, "parking.query", "Mock·演示数据", "✓", "剩余135个车位，并返回P1、P2、V1区域明细。", "确定性Mock数据。"),
    (17, "restaurant.find", "Mock·演示数据", "✓", "返回4处餐饮点，并按拥挤度排序。", "确定性Mock数据。"),
    (18, "meeting.query", "Mock·演示数据", "✓", "返回4间可用会议室及对应空闲时段。", "确定性Mock数据。"),
    (19, "meeting.reserve", "Mock·演示数据", "✗", "未调用工具，模型直接追问楼栋、人数和时间。", "未触发：槽位不完整时，模型选择先澄清必要信息。"),
]

WIDTHS = [Cm(0.8), Cm(4.0), Cm(3.4), Cm(1.4), Cm(8.4), Cm(9.0)]
HEADERS = ["#", "能力", "后端类型", "单轮触发", "真实结果", "验证结论"]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:shd"))
    if old is not None:
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(cell, top=65, start=80, bottom=65, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:cantSplit")
    tr_pr.append(tag)


def set_cell_text(cell, text, *, bold=False, color="1F1F1F", size=7.8, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(text))
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


doc = Document(SOURCE)

section = doc.add_section(WD_SECTION.NEW_PAGE)
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Mm(297)
section.page_height = Mm(210)
section.top_margin = Cm(1.3)
section.bottom_margin = Cm(1.3)
section.left_margin = Cm(1.3)
section.right_margin = Cm(1.3)
section.header_distance = Cm(0.6)
section.footer_distance = Cm(0.6)

heading = doc.add_heading("四、能力验证结果汇总", level=1)
heading.paragraph_format.space_before = Pt(0)
heading.paragraph_format.space_after = Pt(7)

summary = doc.add_paragraph()
summary.paragraph_format.space_after = Pt(9)
summary.paragraph_format.line_spacing = 1.15
summary_run = summary.add_run(
    "验证概览：共覆盖19项能力，其中17项可单轮触发，2项未触发；真实后端调用、前端指令、本地能力与Mock演示数据均已分层标识。"
)
summary_run.font.name = "微软雅黑"
summary_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
summary_run.font.size = Pt(9.5)
summary_run.font.bold = True
summary_run.font.color.rgb = RGBColor(47, 85, 151)

table = doc.add_table(rows=1, cols=len(HEADERS))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.style = "Table Grid"

header = table.rows[0]
repeat_header(header)
keep_row(header)
for index, (cell, text, width) in enumerate(zip(header.cells, HEADERS, WIDTHS)):
    cell.width = width
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    margins(cell, 90, 75, 90, 75)
    shade(cell, "2F5597")
    set_cell_text(cell, text, bold=True, color="FFFFFF", size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)

attention_rows = {2, 4, 7, 8, 10, 19}
for number, capability, backend, triggered, result, conclusion in ROWS:
    row = table.add_row()
    keep_row(row)
    values = [number, capability, backend, triggered, result, conclusion]
    for index, (cell, text, width) in enumerate(zip(row.cells, values, WIDTHS)):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        margins(cell)
        base_fill = "F7F9FC" if number % 2 == 0 else "FFFFFF"
        shade(cell, base_fill)
        align = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
        color = "1F1F1F"
        bold = index in (0, 1)
        if index == 3:
            color = "548235" if triggered == "✓" else "C65911"
            bold = True
        if index == 2:
            backend_fill = (
                "E2F0D9" if backend.startswith("真实") else
                "DDEBF7" if backend.startswith("前端") else
                "E7E6E6" if backend.startswith("本地") else
                "FFF2CC"
            )
            shade(cell, backend_fill)
        if index == 5 and number in attention_rows:
            shade(cell, "FCE4D6")
            color = "9C5700"
        set_cell_text(cell, text, bold=bold, color=color, align=align)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(7)
note.paragraph_format.space_after = Pt(0)
note_run = note.add_run(
    "注：✓表示单轮内已触发目标能力；✗表示未触发目标工具或流程。needs_confirmation 表示已生成待确认草稿或票据，尚未执行实际变更。"
)
note_run.font.name = "微软雅黑"
note_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
note_run.font.size = Pt(8)
note_run.font.color.rgb = RGBColor(89, 89, 89)

doc.core_properties.subject = "设备状态、工单管理与能力验证结果"
doc.save(OUTPUT)
print(OUTPUT)
