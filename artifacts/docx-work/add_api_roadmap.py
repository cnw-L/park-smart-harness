from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(r"D:\park-smart-harness")
SOURCE = ROOT / "artifacts" / "状况-排版优化及能力验证版.docx"
OUTPUT = ROOT / "artifacts" / "状况-排版优化及能力规划版.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:shd"))
    if old is not None:
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_margins(cell, value=95):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row(row, repeat=False):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_text(cell, text, *, bold=False, color="1F1F1F", size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(str(text))
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, headers, rows, widths, *, accent_rows=None):
    accent_rows = accent_rows or set()
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    keep_row(header, repeat=True)
    for cell, text, width in zip(header.cells, headers, widths):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_margins(cell)
        shade(cell, "2F5597")
        set_text(cell, text, bold=True, color="FFFFFF", size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, values in enumerate(rows):
        row = table.add_row()
        keep_row(row)
        for col, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell)
            fill = "F7F9FC" if idx % 2 else "FFFFFF"
            if idx in accent_rows and col == len(values) - 1:
                fill = "FCE4D6"
            shade(cell, fill)
            set_text(cell, text, bold=col == 0, color="9C5700" if idx in accent_rows and col == len(values) - 1 else "1F1F1F")
    return table


def heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    return p


def body(doc, text, *, bold=False, color="404040", after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return p


doc = Document(SOURCE)
section = doc.add_section(WD_SECTION.NEW_PAGE)
section.orientation = WD_ORIENT.PORTRAIT
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(1.7)
section.right_margin = Cm(1.7)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

heading(doc, "五、后端 API 能力差距与建设规划", 1)
body(
    doc,
    "平台后端 API 目录共覆盖58个模块、约450个接口。当前智能助手仅接入设备、健康、工单、告警、事件、能耗和菜单导航等基础域，已验证通用接入架构可复用，但仍有大量物业运营能力尚未暴露给自然语言交互。",
    bold=True,
    color="2F5597",
    after=10,
)

heading(doc, "5.1 当前能力基座", 2)
current_rows = [
    ("设备运行", "device.status / control / health", "设备网关、健康网关", "已接入"),
    ("工单管理", "work_order.query / create", "工单、报修模块", "已接入"),
    ("告警管理", "alarm.query / create", "告警记录模块", "部分接入"),
    ("事件管理", "event.query / create", "事件管理模块", "已接入"),
    ("能耗管理", "energy.query", "能耗概览模块", "仅总量/成本，能力偏弱"),
    ("界面导航", "ui.navigate", "菜单树解析", "已接入"),
    ("知识问答", "knowledge.qa", "Milvus + rerank", "已接入，严格接地"),
]
add_table(doc, ["领域", "已暴露能力", "后端依据", "覆盖情况"], current_rows, [Cm(3.0), Cm(5.0), Cm(4.2), Cm(4.8)], accent_rows={2, 4})

body(
    doc,
    "接入模式已经跑通：每项新能力由 gateway 完成真实接口适配，以 SchemaCapability 注册至 build_default_registry，再由 react 循环自动暴露；所有写操作继续统一走 needs_confirmation 确认门控。",
    bold=True,
    color="548235",
    after=8,
)

heading(doc, "5.2 可扩展能力清单", 2)
expansion_rows = [
    ("A｜运维规划", "巡检计划查询", "巡检计划分页、获取、新增、停用", "今天有哪些巡检任务？", "只读优先；创建/修改需确认", "P0"),
    ("A｜运维规划", "维保计划查询", "维保计划及按设备类型查询", "下月哪些维保到期？", "可与设备健康联动", "P0"),
    ("B｜能耗深化", "多维能耗分析", "分项/区域占比、设备排名、单位面积能耗、碳排放、房间用电", "哪栋楼最耗电？", "补强现有 energy.query", "P0"),
    ("C｜资产财务", "租户与房间查询", "租户账号、名称查询、绑定房间、通行楼层", "A栋301租户是谁？", "只读", "P1"),
    ("C｜资产财务", "合同查询", "按租户查询、到期数量、绑定房间", "本月多少合同到期？", "只读", "P1"),
    ("C｜资产财务", "账单与缴费查询", "账单信息、缴费及支付记录", "XX租户本月欠费多少？", "只读；支付操作需确认", "P1"),
    ("C｜资产财务", "发票申请", "发票抬头、发票申请", "帮我申请本月发票。", "写操作，必须确认", "P1"),
    ("D｜告警闭环", "告警升级/误报", "告警升级、误报标记、点位设备参数", "把这条告警标记为误报。", "写操作，必须确认", "P1"),
    ("E｜通知知识", "通知公告查询", "通知分页、搜索、已读状态", "最近有什么园区公告？", "可同步进入知识库", "P1"),
    ("F｜场景联动", "自然语言场景执行", "场景动作、可控设备树、场景新增", "下班后关闭公共区照明。", "差异化高、复杂度高", "P2"),
    ("F｜位置安全", "定位与电子围栏", "实时/历史定位、围栏告警", "XX定位器现在在哪里？", "依赖定位硬件部署", "P2"),
    ("F｜通行控制", "电梯与楼层权限", "楼层配置、租户通行楼层", "给该租户开放10楼权限。", "权限写入需确认和审计", "P2"),
    ("F｜便捷能力", "天气与点位收藏", "高德天气、用户点位收藏、保存点位", "今天园区天气怎么样？", "低成本增强", "P2"),
]
add_table(
    doc,
    ["分层", "能力", "接口依据", "示例问法", "性质/约束", "优先级"],
    expansion_rows,
    [Cm(2.1), Cm(2.9), Cm(4.4), Cm(3.4), Cm(3.5), Cm(1.2)],
    accent_rows={9, 10, 11},
)

heading(doc, "5.3 Mock 能力升级边界", 2)
mock_rows = [
    ("parking.query", "可部分升级", "后端已有停车缴费订单与月卡接口，但并不等同于实时剩余车位。", "拆分为“停车缴费/月卡查询”；剩余车位仍保留Mock或另接停车场实时系统。"),
    ("restaurant.find", "暂保留Mock", "API目录未发现匹配的餐饮点与拥挤度数据源。", "待接入商户或客流系统后再真实化。"),
    ("meeting.query/reserve", "暂保留Mock", "API目录未发现会议室资源与预订接口。", "保留槽位澄清逻辑，后续对接会议室系统。"),
]
add_table(doc, ["现有能力", "结论", "接口差距", "建议"], mock_rows, [Cm(3.4), Cm(2.4), Cm(5.4), Cm(6.3)], accent_rows={0})

heading(doc, "5.4 建议实施路线", 2)
roadmap_rows = [
    ("P0｜第一批", "巡检计划、维保计划、能耗多维分析", "只读、高频、风险低，与 device.health / energy 直接联动。", "优先完成 gateway 映射、槽位Schema和真实数据回归测试。"),
    ("P1｜第二批", "租户、合同、账单、发票、告警闭环、公告", "覆盖物业核心运营查询，并补齐告警处置闭环。", "查询先行；发票、支付、告警修改统一接入确认门控与审计。"),
    ("P2｜第三批", "场景联动、定位围栏、楼层权限、天气、点位收藏", "形成差异化能力，但依赖硬件、权限或跨系统编排。", "按实际部署条件逐项验证，场景执行需强化风险分级。"),
]
add_table(doc, ["阶段", "建设范围", "价值判断", "交付重点"], roadmap_rows, [Cm(2.6), Cm(4.5), Cm(5.2), Cm(5.2)])

heading(doc, "5.5 关键约束与验收重点", 2)
risk_rows = [
    ("数据可用性", "真实调用成功不等于有业务数据；能耗和告警需明确默认时间窗口与空结果语义。"),
    ("能力边界", "停车缴费/月卡不能替代实时车位余量，禁止以相邻接口伪装成同一能力。"),
    ("写操作治理", "创建、修改、支付、权限和场景执行必须使用 needs_confirmation，并记录审计信息。"),
    ("硬件依赖", "定位、围栏、摄像头和设备控制应先核实园区实际部署、在线状态及访问权限。"),
    ("知识接地", "通知公告可进入Milvus，但需保留来源、发布时间和有效期，避免过期内容参与回答。"),
]
add_table(doc, ["关注项", "验收要求"], risk_rows, [Cm(3.2), Cm(14.3)])

body(
    doc,
    "总体建议：先以巡检/维保查询和能耗分析深化建立第一批真实能力，再扩展租户、合同等物业核心域；场景联动作为差异化能力压轴实施。",
    bold=True,
    color="2F5597",
    after=0,
)

doc.core_properties.subject = "设备状态、能力验证与后端API建设规划"
doc.save(OUTPUT)
print(OUTPUT)
