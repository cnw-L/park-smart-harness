from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\park-smart-harness")
SOURCE = ROOT / "artifacts" / "状况-最终整理版.docx"
OUTPUT = ROOT / "artifacts" / "状况-最终整理版（第四章含中文工具名）.docx"

CHINESE_NAMES = {
    "device.status.query": "设备状态查询",
    "device.control_ticket.prepare": "设备控制票据准备",
    "device.health.query": "设备健康查询",
    "energy.query": "能耗查询",
    "work_order.query": "工单查询",
    "work_order.create_prepare": "工单创建准备",
    "alarm.query": "告警查询",
    "alarm.create_order": "告警建单",
    "event.query": "事件查询",
    "event.create_prepare": "事件上报准备",
    "ui.navigate": "页面导航",
    "knowledge.qa": "知识问答",
    "ui.focus_camera": "摄像头聚焦",
    "ui.open_device_panel": "打开设备面板",
    "date.current": "当前日期查询",
    "parking.query": "停车位查询",
    "restaurant.find": "餐饮点查询",
    "meeting.query": "会议室查询",
    "meeting.reserve": "会议室预订",
}

HEADERS = ["#", "能力", "中文工具名", "后端类型", "单轮触发", "真实结果", "验证结论"]
WIDTHS = [Cm(0.7), Cm(3.6), Cm(2.5), Cm(3.3), Cm(1.3), Cm(7.7), Cm(8.0)]
ATTENTION_ROWS = {2, 4, 7, 8, 10, 19}


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(cell, value=60):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for tag in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def keep_row(row, repeat=False):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_text(cell, text, *, bold=False, color="1F1F1F", align=WD_ALIGN_PARAGRAPH.LEFT, size=7.3):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


doc = Document(SOURCE)
old_table = next(
    table for table in doc.tables
    if table.rows and [cell.text for cell in table.rows[0].cells][:2] == ["#", "能力"]
)

old_rows = [[cell.text for cell in row.cells] for row in old_table.rows[1:]]

new_table = doc.add_table(rows=1, cols=7)
new_table.alignment = WD_TABLE_ALIGNMENT.CENTER
new_table.autofit = False
new_table.style = "Table Grid"

header = new_table.rows[0]
keep_row(header, repeat=True)
for cell, text, width in zip(header.cells, HEADERS, WIDTHS):
    cell.width = width
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    margins(cell)
    shade(cell, "2F5597")
    set_text(cell, text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.8)

for old_values in old_rows:
    number = int(old_values[0])
    capability = old_values[1]
    values = [old_values[0], capability, CHINESE_NAMES[capability], *old_values[2:]]
    row = new_table.add_row()
    keep_row(row)
    for col, (cell, text, width) in enumerate(zip(row.cells, values, WIDTHS)):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        margins(cell)
        fill = "F7F9FC" if number % 2 == 0 else "FFFFFF"
        if col == 2:
            fill = "EAF2F8"
        if col == 3:
            fill = (
                "E2F0D9" if text.startswith("真实") else
                "DDEBF7" if text.startswith("前端") else
                "E7E6E6" if text.startswith("本地") else
                "FFF2CC"
            )
        if col == 6 and number in ATTENTION_ROWS:
            fill = "FCE4D6"
        shade(cell, fill)
        color = "1F1F1F"
        if col == 4:
            color = "548235" if text == "✓" else "C65911"
        if col == 6 and number in ATTENTION_ROWS:
            color = "9C5700"
        set_text(
            cell,
            text,
            bold=col in (0, 1, 2, 4),
            color=color,
            align=WD_ALIGN_PARAGRAPH.CENTER if col in (0, 4) else WD_ALIGN_PARAGRAPH.LEFT,
        )

old_table._tbl.addprevious(new_table._tbl)
old_table._element.getparent().remove(old_table._element)

doc.save(OUTPUT)
print(OUTPUT)
