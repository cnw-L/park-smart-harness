from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(r"D:\park-smart-harness")
SOURCE = ROOT / "artifacts" / "状况-简化能力规划版.docx"
OUTPUT = ROOT / "artifacts" / "状况-工具优化说明版.docx"

ISSUES = [
    ("参数重复识别", "所有工具", "一次调用会重复识别用户参数，可能增加响应时间，并出现前后结果不一致。", "优先使用模型已经填写的参数，缺失时再补充识别。", "中"),
    ("时间格式不统一", "告警、能耗、事件、工单、会议室查询", "用户说“今天”或“近7天”时，工具内部接收的时间格式不一致。", "统一时间参数格式，并保留自然语言时间的自动转换。", "高"),
    ("部分操作依赖上一轮", "告警建单、会议室预订等", "没有上一轮查询结果或必要信息时，工具无法直接执行。", "缺少信息时明确提示；适合的场景可先自动查询，再让用户确认。", "高"),
    ("确认流程较独立", "设备控制、工单创建、事件上报、告警建单", "准备、确认和执行分成不同步骤，状态管理较复杂。", "继续保留确认机制，同时让待确认内容和执行结果更清楚。", "中"),
    ("设备名称容易误匹配", "设备状态、摄像头聚焦、设备面板", "设备名称中的数字、空格或房间号可能被误当成设备ID。", "统一清理空格，区分设备名称和设备ID，并增加匹配校验。", "高"),
    ("事件标题不够干净", "event.create_prepare", "“我要上报一个事件”等对话前缀会被带入标题和描述。", "创建草稿前清理口语前缀，只保留真正的事件内容。", "立即"),
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(cell, value=95):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for tag in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def keep_row(row, repeat=False):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_text(cell, text, *, bold=False, color="1F1F1F", align=WD_ALIGN_PARAGRAPH.LEFT, size=8.7):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def body(doc, text, *, bold=False, color="404040", after=7):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10.5)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return p


doc = Document(SOURCE)
section = doc.add_section(WD_SECTION.NEW_PAGE)
section.orientation = WD_ORIENT.PORTRAIT
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(1.6)
section.right_margin = Cm(1.6)

heading = doc.add_heading("六、现有工具问题与后续优化", level=1)
heading.paragraph_format.space_before = Pt(0)
heading.paragraph_format.space_after = Pt(8)

body(
    doc,
    "系统共注册19项能力，其中17项可以单轮触发；告警建单和会议室预订需要上一轮结果或补充信息。当前流程支持多步调用，主要问题集中在参数格式、上下文衔接和确认流程。",
    bold=True,
    color="2F5597",
    after=9,
)

headers = ["问题", "影响范围", "主要表现", "优化方式", "优先级"]
widths = [Cm(2.5), Cm(3.6), Cm(4.5), Cm(5.0), Cm(1.5)]
table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.style = "Table Grid"

header = table.rows[0]
keep_row(header, repeat=True)
for cell, text, width in zip(header.cells, headers, widths):
    cell.width = width
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    margins(cell)
    shade(cell, "2F5597")
    set_text(cell, text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, size=8.9)

for index, values in enumerate(ISSUES):
    row = table.add_row()
    keep_row(row)
    for col, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        margins(cell)
        fill = "F7F9FC" if index % 2 else "FFFFFF"
        if col == 4:
            fill = "FCE4D6" if text in ("立即", "高") else "FFF2CC"
        shade(cell, fill)
        color = "C65911" if col == 4 and text in ("立即", "高") else "1F1F1F"
        set_text(
            cell,
            text,
            bold=col in (0, 4),
            color=color,
            align=WD_ALIGN_PARAGRAPH.CENTER if col == 4 else WD_ALIGN_PARAGRAPH.LEFT,
        )

sub = doc.add_heading("建议优化顺序", level=2)
sub.paragraph_format.space_before = Pt(12)
sub.paragraph_format.space_after = Pt(5)

steps = [
    ("第一步", "清理事件标题，完善设备名称匹配。", "改动小、风险低，可以直接改善用户看到的结果。"),
    ("第二步", "统一时间参数，补强缺少上下文时的提示或自动查询。", "解决查询不稳定和部分工具单轮难触发的问题。"),
    ("第三步", "减少重复参数识别，整理确认与执行状态。", "降低响应时间，并让整体流程更稳定。"),
]
step_table = doc.add_table(rows=1, cols=3)
step_table.alignment = WD_TABLE_ALIGNMENT.CENTER
step_table.autofit = False
step_table.style = "Table Grid"
step_widths = [Cm(2.2), Cm(6.6), Cm(8.3)]
step_header = step_table.rows[0]
keep_row(step_header, repeat=True)
for cell, text, width in zip(step_header.cells, ["阶段", "优化内容", "目标"], step_widths):
    cell.width = width
    margins(cell)
    shade(cell, "5B9BD5")
    set_text(cell, text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
for index, values in enumerate(steps):
    row = step_table.add_row()
    keep_row(row)
    for col, (cell, text, width) in enumerate(zip(row.cells, values, step_widths)):
        cell.width = width
        margins(cell)
        shade(cell, "FFFFFF" if index % 2 == 0 else "F7F9FC")
        set_text(cell, text, bold=col == 0, align=WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.LEFT, size=9.2)

body(
    doc,
    "说明：这些问题不影响当前基本演示，但会影响响应速度、复杂问法和连续操作的稳定性，因此建议按以上顺序逐步优化。",
    color="595959",
    after=0,
)

doc.core_properties.subject = "设备状态、能力验证、简单能力规划与工具优化说明"
doc.save(OUTPUT)
print(OUTPUT)
