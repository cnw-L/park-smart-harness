from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(r"D:\park-smart-harness")
SOURCE = ROOT / "artifacts" / "状况-排版优化及能力验证版.docx"
OUTPUT = ROOT / "artifacts" / "状况-简化能力规划版.docx"

TOOLS = [
    ("巡检计划查询", "查询今天、本周有哪些巡检任务。", "今天有哪些巡检任务？", "优先增加"),
    ("维保计划查询", "查询设备什么时候需要保养或维修。", "空调机组下次什么时候维保？", "优先增加"),
    ("能耗排行查询", "查看楼栋、区域或设备的能耗排行。", "本月哪栋楼最耗电？", "优先增加"),
    ("租户信息查询", "根据房间或公司名称查询租户信息。", "A栋301的租户是谁？", "后续增加"),
    ("合同到期查询", "查询合同到期时间和近期到期数量。", "本月有多少合同到期？", "后续增加"),
    ("园区公告查询", "查询最新通知和园区公告。", "最近有什么园区公告？", "后续增加"),
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for tag, value in (("top", 110), ("start", 120), ("bottom", 110), ("end", 120)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def keep_row(row, repeat=False):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_text(cell, text, *, bold=False, color="1F1F1F", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.18
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9.5)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, *, bold=False, color="404040"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10.5)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return p


doc = Document(SOURCE)
section = doc.add_section(WD_SECTION.NEW_PAGE)
section.orientation = WD_ORIENT.PORTRAIT
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

heading = doc.add_heading("五、可增加的简单工具能力", level=1)
heading.paragraph_format.space_before = Pt(0)
heading.paragraph_format.space_after = Pt(8)

add_body(
    doc,
    "目前系统已经支持设备状态、设备健康、工单、告警、事件、能耗和页面导航等功能。后续不需要一次增加太多工具，可以先从日常使用频率高、容易理解的查询功能开始。",
    bold=True,
    color="2F5597",
)

headers = ["建议工具", "能做什么", "示例问法", "安排"]
widths = [Cm(3.4), Cm(5.0), Cm(5.4), Cm(2.6)]
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.style = "Table Grid"

header = table.rows[0]
keep_row(header, repeat=True)
for cell, text, width in zip(header.cells, headers, widths):
    cell.width = width
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    margins(cell)
    shade(cell, "2F5597")
    set_text(cell, text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)

for index, values in enumerate(TOOLS):
    row = table.add_row()
    keep_row(row)
    for col, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        margins(cell)
        shade(cell, "F7F9FC" if index % 2 else "FFFFFF")
        if col == 3:
            shade(cell, "E2F0D9" if text == "优先增加" else "FFF2CC")
        set_text(
            cell,
            text,
            bold=col in (0, 3),
            color="548235" if col == 3 and text == "优先增加" else "1F1F1F",
            align=WD_ALIGN_PARAGRAPH.CENTER if col == 3 else WD_ALIGN_PARAGRAPH.LEFT,
        )

subheading = doc.add_heading("建议先做", level=2)
subheading.paragraph_format.space_before = Pt(12)
subheading.paragraph_format.space_after = Pt(5)
add_body(doc, "第一步先增加巡检计划、维保计划和能耗排行三个查询工具。这三项都是只读查询，使用频率高，也不会直接修改系统数据。")

subheading = doc.add_heading("后续再做", level=2)
subheading.paragraph_format.space_before = Pt(8)
subheading.paragraph_format.space_after = Pt(5)
add_body(doc, "租户信息、合同到期和园区公告可以作为第二批能力，等第一批运行稳定后再逐步加入。")

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(5)
note.paragraph_format.space_after = Pt(0)
note_run = note.add_run("说明：如果以后增加创建、修改或控制类功能，执行前仍需由用户确认。")
note_run.font.name = "微软雅黑"
note_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
note_run.font.size = Pt(9)
note_run.font.color.rgb = RGBColor(89, 89, 89)

doc.core_properties.subject = "设备状态、能力验证与简化能力规划"
doc.save(OUTPUT)
print(OUTPUT)
